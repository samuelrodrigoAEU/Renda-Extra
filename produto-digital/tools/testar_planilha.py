# Preenche uma cópia com o caso do brigadeiro, recalcula no LibreOffice e confere com a conta feita à mão.
import sys, subprocess, shutil, os
from openpyxl import load_workbook
SRC, TMP = sys.argv[1], sys.argv[2]
os.makedirs(TMP, exist_ok=True)
wb = load_workbook(SRC)
c = wb['Custos do mês']; c['B3'] = 2000; c['B4'] = 120; c['B7'] = 120; c['B8'] = 150; c['B11'] = 90   # fixos = 360 → 3/h
i = wb['Insumos']
precos = {'Leite condensado': 7, 'Manteiga sem sal': 9, 'Chocolate em pó 50%': 8, 'Granulado': 12}
for r in range(5, 105):
    if i[f'A{r}'].value in precos: i[f'D{r}'] = precos[i[f'A{r}'].value]
f = wb['Ficha 1']
f['B2'] = 'Brigadeiro'; f['B3'] = 50; f['B4'] = 90; f['B5'] = 0.06; f['B36'] = 1.5
for r, (n, q) in enumerate([('Leite condensado', 790), ('Manteiga sem sal', 20), ('Chocolate em pó 50%', 40), ('Granulado', 120), ('Coco ralado', 10)], start=8):
    f[f'A{r}'] = n; f[f'B{r}'] = q
wb.save(f'{TMP}/in.xlsx')
subprocess.run(['soffice', '--headless', '--calc', '--convert-to', 'xlsx:Calc MS Excel 2007 XML', '--outdir', f'{TMP}/out', f'{TMP}/in.xlsx'], check=True, capture_output=True)
v = load_workbook(f'{TMP}/out/in.xlsx', data_only=True)
F, E, T, C = v['Ficha 1'], v['Exemplo'], v['Tabela WhatsApp'], v['Custos do mês']
custo = (7/395*790 + 9/200*20 + 8/200*40 + 12/500*120)/50 + 0.06 + 90/60*(2000/120)/50 + 90/60*3/50
import math
esp = {'custo': custo, 'min': math.ceil(custo/0.96*2)/2, 'sug': math.ceil(custo/0.66*2)/2}
res = []
ok = lambda n, c: res.append(('PASS' if c else 'FAIL') + ' ' + n)
ok(f"valor da hora {C['B5'].value:.2f}", abs(C['B5'].value - 16.6667) < 0.01)
ok(f"fixo por hora {C['B15'].value:.2f}", abs(C['B15'].value - 3) < 1e-6)
ok(f"custo real {F['D29'].value:.4f} = {esp['custo']:.4f}", abs(F['D29'].value - esp['custo']) < 1e-4)
ok(f"preço mínimo {F['D31'].value} = {esp['min']}", F['D31'].value == esp['min'])
ok(f"preço sugerido {F['D32'].value} = {esp['sug']}", F['D32'].value == esp['sug'])
ok(f"insumo sem preço avisa ({F['D12'].value})", F['D12'].value == 'preencha o preço')
ok(f"preço de hoje R$1,50 mostra sobra {F['D37'].value:.3f}", abs(F['D37'].value - (1.5*0.96 - esp['custo'])) < 1e-4)
ok(f"app vazio sem taxa ({F['D34'].value!r})", F['D34'].value in (None, ''))
ok(f"tabela WhatsApp: {T['D5'].value!r}", T['D5'].value and T['D5'].value.startswith('Brigadeiro: R$ 2'))
ok(f"ficha vazia não polui a tabela ({T['A6'].value!r})", T['A6'].value in (None, ''))
ok(f"exemplo = ficha preenchida ({E['B23'].value} / {F['D32'].value})", E['B23'].value == F['D32'].value)
print('\n'.join(res)); sys.exit(any(r.startswith('FAIL') for r in res))
