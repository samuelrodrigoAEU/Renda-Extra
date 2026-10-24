# Gera "Conta Doce" (planilha de precificação para confeitaria) e o order bump "Agenda de Encomendas".
# Uso: PYTHONPATH=<openpyxl> python3 gerar_planilha.py <pasta_saida>
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule

OUT = sys.argv[1]
ROSA, CREME, TINTA, CINZA, VERDE = 'F6D7D2', 'FFF8EF', '2B2420', '7A6E66', 'DDEFD9'
fill = lambda c: PatternFill('solid', fgColor=c)
H1 = Font(name='Arial', size=16, bold=True, color=TINTA)
B = Font(name='Arial', size=11, bold=True, color=TINTA)
N = Font(name='Arial', size=11, color=TINTA)
S = Font(name='Arial', size=10, italic=True, color=CINZA)
INP = fill(CREME)        # célula para a confeiteira preencher
RES = fill(VERDE)        # resultado calculado
fino = Side(style='thin', color='E5D9CF')
caixa = Border(left=fino, right=fino, top=fino, bottom=fino)
R = '"R$" #,##0.00'
PCT = '0%'

def cell(ws, ref, v=None, font=N, fmt=None, f=None, wrap=False):
    c = ws[ref]
    if v is not None: c.value = v
    c.font = font
    if fmt: c.number_format = fmt
    if f: c.fill = f; c.border = caixa
    c.alignment = Alignment(wrap_text=wrap, vertical='center')
    return c

def larguras(ws, w):
    for col, x in w.items(): ws.column_dimensions[col].width = x

wb = Workbook()

# 1. Comece aqui
ws = wb.active; ws.title = 'Comece aqui'
larguras(ws, {'A': 4, 'B': 92})
cell(ws, 'B2', 'Conta Doce · o preço do seu doce em 15 minutos', H1)
passos = [
    ('Células em creme', 'são as que você preenche. As verdes são calculadas sozinhas. Não apague fórmulas.'),
    ('Passo 1 · Custos do mês (5 min)', 'quanto você quer ganhar por mês, quantas horas trabalha na confeitaria e os gastos fixos (gás, luz, internet, DAS).'),
    ('Passo 2 · Insumos (5 min)', 'escreva o preço que você pagou só nos ingredientes do seu doce mais vendido. A lista já traz tamanhos comuns de embalagem: confira se o seu é igual.'),
    ('Passo 3 · Ficha 1 (5 min)', 'escolha os ingredientes na lista, diga quanto usa, quantas unidades rende e quanto tempo leva. Pronto: aparece o custo real, o preço mínimo e o preço sugerido.'),
    ('Tabela WhatsApp', 'junta os preços de todas as fichas em linhas prontas para copiar e mandar para clientes.'),
    ('Exemplo', 'mostra uma ficha preenchida com valores de exemplo. Os valores do exemplo NÃO são preços de mercado.'),
    ('Limites', 'a planilha calcula a partir dos números que você coloca. Ela não sabe o preço da sua região nem garante vendas. Revise os preços dos insumos quando fizer compras.'),
]
r = 4
for t, d in passos:
    cell(ws, f'B{r}', t, B); cell(ws, f'B{r+1}', d, N, wrap=True); ws.row_dimensions[r+1].height = 32; r += 3

# 2. Custos do mês
cm = wb.create_sheet('Custos do mês')
larguras(cm, {'A': 46, 'B': 18, 'C': 60})
cell(cm, 'A1', 'Custos do mês', H1)
linhas = [
    (3, 'Quanto você quer ganhar por mês pelo seu trabalho', 2000, R, 'seu pró-labore: o salário que você se paga'),
    (4, 'Horas por mês trabalhando na confeitaria', 120, '0', 'conte produção, compras, embalagem e atendimento'),
    (7, 'Gás (parte da confeitaria)', 0, R, ''), (8, 'Energia (parte da confeitaria)', 0, R, ''),
    (9, 'Água (parte da confeitaria)', 0, R, ''), (10, 'Internet e celular (parte do negócio)', 0, R, ''),
    (11, 'DAS do MEI ou impostos', 0, R, 'confira o valor do seu boleto'),
    (12, 'Utensílios, manutenção e reposição', 0, R, ''), (13, 'Outros gastos fixos', 0, R, ''),
    (17, 'Taxa da maquininha ou do link de pagamento', 0.04, PCT, 'exemplo de 4%: troque pela taxa do seu contrato'),
    (18, 'Taxa do aplicativo de delivery (se usar)', 0, PCT, 'confira no seu contrato; deixe 0% se não usa'),
    (19, 'Margem de lucro desejada', 0.3, PCT, 'quanto sobra para o negócio além do seu pró-labore'),
]
for row, t, v, fmt, nota in linhas:
    cell(cm, f'A{row}', t); cell(cm, f'B{row}', v, fmt=fmt, f=INP); cell(cm, f'C{row}', nota, S)
cell(cm, 'A6', 'Gastos fixos por mês', B)
cell(cm, 'A5', 'Valor da sua hora', B); cell(cm, 'B5', '=IFERROR(B3/B4,0)', B, R, RES)
cell(cm, 'A14', 'Total de gastos fixos', B); cell(cm, 'B14', '=SUM(B7:B13)', B, R, RES)
cell(cm, 'A15', 'Gasto fixo por hora trabalhada', B); cell(cm, 'B15', '=IFERROR(B14/B4,0)', B, R, RES)
cell(cm, 'A16', 'Taxas e margem', B)

# 3. Insumos
ins = wb.create_sheet('Insumos')
larguras(ins, {'A': 34, 'B': 10, 'C': 16, 'D': 16, 'E': 18})
cell(ins, 'A1', 'Insumos e embalagens', H1)
cell(ins, 'A2', 'Preencha o preço só do que você usa. Tamanhos de embalagem são comuns no mercado: confira o seu.', S)
for col, t in zip('ABCDE', ['Insumo', 'Unidade', 'Qtd na embalagem', 'Preço pago', 'Custo por unidade']):
    cell(ins, f'{col}4', t, B)
itens = [
 ('Leite condensado','g',395),('Creme de leite','g',200),('Leite integral','ml',1000),('Leite em pó','g',200),
 ('Manteiga sem sal','g',200),('Margarina','g',500),('Açúcar refinado','g',1000),('Açúcar de confeiteiro','g',500),
 ('Farinha de trigo','g',1000),('Fermento químico','g',100),('Amido de milho','g',500),('Ovos','un',12),
 ('Chocolate em pó 50%','g',200),('Cacau em pó 100%','g',200),('Chocolate meio amargo (barra)','g',1000),
 ('Chocolate ao leite (barra)','g',1000),('Chocolate branco (barra)','g',1000),('Granulado','g',500),
 ('Coco ralado','g',100),('Leite de coco','ml',200),('Doce de leite','g',400),('Nutella ou creme de avelã','g',350),
 ('Morango','g',250),('Limão','un',1),('Maracujá (polpa)','g',100),('Essência de baunilha','ml',30),
 ('Óleo','ml',900),('Glucose de milho','g',350),('Gelatina sem sabor','g',12),('Corante em gel','g',15),
 ('Paçoca','un',50),('Castanha de caju','g',100),('Nozes','g',100),('Amendoim torrado','g',500),
 ('Biscoito maisena','g',200),('Chantilly (pó ou pronto)','g',500),('Glacê pronto','g',500),('Pasta americana','g',800),
 ('Forminha nº 4','un',100),('Forminha nº 5','un',100),('Caixa para 4 docinhos','un',1),('Caixa para 9 docinhos','un',1),
 ('Pote 250 ml com tampa','un',1),('Pote 120 ml com tampa','un',1),('Saquinho celofane','un',100),('Fita de cetim (rolo)','m',10),
 ('Etiqueta adesiva','un',100),('Embalagem para fatia de bolo','un',1),('Caixa de bolo 20 cm','un',1),('Base de papelão para bolo','un',1),
 ('Colher descartável','un',50),('Sacola','un',1),
]
for i, (n, u, q) in enumerate(itens):
    r = 5 + i
    cell(ins, f'A{r}', n, f=INP); cell(ins, f'B{r}', u, f=INP); cell(ins, f'C{r}', q, f=INP); cell(ins, f'D{r}', None, fmt=R, f=INP)
    cell(ins, f'E{r}', f'=IF(AND(C{r}>0,D{r}>0),D{r}/C{r},"")', fmt='"R$" #,##0.0000', f=RES)
for r in range(5 + len(itens), 105):   # linhas livres para insumos novos
    for col in 'ABCD': cell(ins, f'{col}{r}', None, fmt=R if col == 'D' else None, f=INP)
    cell(ins, f'E{r}', f'=IF(AND(C{r}>0,D{r}>0),D{r}/C{r},"")', fmt='"R$" #,##0.0000', f=RES)

# 4. Fichas
def ficha(nome):
    fs = wb.create_sheet(nome)
    larguras(fs, {'A': 42, 'B': 14, 'C': 10, 'D': 19, 'E': 44})
    cell(fs, 'A1', f'{nome} · ficha de preço', H1)
    cell(fs, 'A2', 'Produto'); cell(fs, 'B2', None, B, f=INP); fs.merge_cells('B2:D2')
    cell(fs, 'A3', 'Rendimento da receita (unidades)'); cell(fs, 'B3', None, fmt='0', f=INP)
    cell(fs, 'A4', 'Tempo para fazer a receita inteira (min)'); cell(fs, 'B4', None, fmt='0', f=INP)
    cell(fs, 'A5', 'Embalagem por unidade vendida'); cell(fs, 'B5', None, fmt=R, f=INP)
    cell(fs, 'E5', 'forminha, caixa, pote, etiqueta: some o que vai em cada unidade', S)
    for col, t in zip('ABCD', ['Ingrediente', 'Qtd usada', 'Unidade', 'Custo']): cell(fs, f'{col}7', t, B)
    dv = DataValidation(type='list', formula1='=Insumos!$A$5:$A$104', allow_blank=True); fs.add_data_validation(dv)
    for r in range(8, 23):
        cell(fs, f'A{r}', None, f=INP); dv.add(f'A{r}')
        cell(fs, f'B{r}', None, fmt='0.##', f=INP)
        cell(fs, f'C{r}', f'=IF(A{r}="","",IFERROR(INDEX(Insumos!$B$5:$B$104,MATCH(A{r},Insumos!$A$5:$A$104,0)),""))', f=RES)
        cell(fs, f'D{r}', f'=IF(OR(A{r}="",B{r}=""),"",IFERROR(B{r}*INDEX(Insumos!$E$5:$E$104,MATCH(A{r},Insumos!$A$5:$A$104,0)),"preencha o preço"))', fmt=R, f=RES)
    fs.conditional_formatting.add('D8:D22', FormulaRule(formula=['ISTEXT(D8)'], font=Font(color='B3261E', bold=True)))
    cell(fs, 'A23', 'Ingredientes da receita', B); cell(fs, 'D23', '=SUM(D8:D22)', B, R, RES)
    linhas = [
        (25, 'Ingredientes por unidade', '=IFERROR(D23/B3,0)'),
        (26, 'Embalagem por unidade', '=N(B5)'),
        (27, 'Seu trabalho por unidade', "=IFERROR(B4/60*'Custos do mês'!B5/B3,0)"),
        (28, 'Gastos fixos por unidade', "=IFERROR(B4/60*'Custos do mês'!B15/B3,0)"),
        (29, 'CUSTO REAL POR UNIDADE', '=SUM(D25:D28)'),
        (31, 'Preço mínimo (sem lucro, já com a taxa)', "=IFERROR(ROUNDUP(D29/(1-'Custos do mês'!B17)*2,0)/2,0)"),
        (32, 'PREÇO SUGERIDO (encomenda)', "=IFERROR(ROUNDUP(D29/(1-'Custos do mês'!B17-'Custos do mês'!B19)*2,0)/2,0)"),
        (33, 'Lucro por unidade no preço sugerido', "=D32*(1-'Custos do mês'!B17)-D29"),
        (34, 'Preço sugerido no aplicativo de delivery', "=IF('Custos do mês'!B18=0,\"\",IFERROR(ROUNDUP(D29/(1-'Custos do mês'!B18-'Custos do mês'!B17-'Custos do mês'!B19)*2,0)/2,0))"),
        (36, 'Preço que você cobra hoje (opcional)', None),
        (37, 'Quanto sobra por unidade no preço de hoje', "=IF(B36=\"\",\"\",B36*(1-'Custos do mês'!B17)-D29)"),
    ]
    for row, t, fml in linhas:
        forte = row in (29, 32)
        cell(fs, f'A{row}', t, B if forte else N)
        if fml: cell(fs, f'D{row}', fml, B if forte else N, R, RES)
    cell(fs, 'B36', None, fmt=R, f=INP)
    cell(fs, 'E31', 'abaixo disso você paga para trabalhar', S); cell(fs, 'E32', 'arredondado para cima em R$ 0,50', S)
    cell(fs, 'E37', 'se der negativo, o preço de hoje não cobre o custo real', S)
    fs.conditional_formatting.add('D37', FormulaRule(formula=['AND(ISNUMBER(D37),D37<0)'], font=Font(color='B3261E', bold=True)))

for i in range(1, 7): ficha(f'Ficha {i}')

# 5. Tabela WhatsApp
tw = wb.create_sheet('Tabela WhatsApp')
larguras(tw, {'A': 32, 'B': 16, 'C': 16, 'D': 50})
cell(tw, 'A1', 'Tabela de preços para copiar', H1)
cell(tw, 'A2', 'Copie a coluna D e cole no WhatsApp ou no seu cardápio.', S)
for col, t in zip('ABCD', ['Produto', 'Encomenda', 'App', 'Linha pronta']): cell(tw, f'{col}4', t, B)
for i in range(1, 7):
    r, f = 4 + i, f"'Ficha {i}'"
    cell(tw, f'A{r}', f'=IF({f}!B2="","",{f}!B2)')
    cell(tw, f'B{r}', f'=IF({f}!B2="","",{f}!D32)', fmt=R)
    cell(tw, f'C{r}', f'=IF(OR({f}!B2="",{f}!D34=""),"",{f}!D34)', fmt=R)
    cell(tw, f'D{r}', f'=IF({f}!B2="","",{f}!B2&": R$ "&FIXED({f}!D32,2))')

# 6. Exemplo (autocontido: não depende da aba Insumos)
ex = wb.create_sheet('Exemplo')
larguras(ex, {'A': 30, 'B': 16, 'C': 16, 'D': 14, 'E': 16})
cell(ex, 'A1', 'Exemplo: brigadeiro tradicional', H1)
cell(ex, 'A2', 'VALORES DE EXEMPLO para mostrar o cálculo. Não são preços de mercado nem da sua região.', Font(name='Arial', size=10, bold=True, color='B3261E'))
for col, t in zip('ABCDE', ['Ingrediente', 'Qtd embalagem', 'Preço exemplo', 'Qtd usada', 'Custo']): cell(ex, f'{col}4', t, B)
dados = [('Leite condensado (g)', 395, 7.0, 790), ('Manteiga (g)', 200, 9.0, 20), ('Chocolate em pó 50% (g)', 200, 8.0, 40),
         ('Granulado (g)', 500, 12.0, 120), ('Forminha nº 4 (un)', 100, 6.0, 50)]
for i, (n, q, p, u) in enumerate(dados):
    r = 5 + i
    cell(ex, f'A{r}', n); cell(ex, f'B{r}', q); cell(ex, f'C{r}', p, fmt=R); cell(ex, f'D{r}', u)
    cell(ex, f'E{r}', f'=C{r}/B{r}*D{r}', fmt=R, f=RES)
ex_l = [(11, 'Rendimento (unidades)', 50, '0'), (12, 'Tempo da receita (min)', 90, '0'), (13, 'Valor da hora (exemplo)', 16.67, R),
        (14, 'Gasto fixo por hora (exemplo)', 3.0, R), (15, 'Taxa do pagamento (exemplo)', 0.04, PCT), (16, 'Margem desejada (exemplo)', 0.3, PCT)]
for row, t, v, fmt in ex_l: cell(ex, f'A{row}', t); cell(ex, f'B{row}', v, fmt=fmt)
calc = [(18, 'Ingredientes e forminha por unidade', '=SUM(E5:E9)/B11'), (19, 'Trabalho por unidade', '=B12/60*B13/B11'),
        (20, 'Gastos fixos por unidade', '=B12/60*B14/B11'), (21, 'CUSTO REAL POR UNIDADE', '=SUM(B18:B20)'),
        (22, 'Preço mínimo', '=ROUNDUP(B21/(1-B15)*2,0)/2'), (23, 'PREÇO SUGERIDO', '=ROUNDUP(B21/(1-B15-B16)*2,0)/2'),
        (24, 'Lucro por unidade', '=B23*(1-B15)-B21')]
for row, t, fml in calc:
    cell(ex, f'A{row}', t, B if row in (21, 23) else N); cell(ex, f'B{row}', fml, B, R, RES)
cell(ex, 'A26', 'Repare: sem contar o trabalho e os gastos fixos, o custo parece bem menor do que é.', S)

for w in wb.worksheets:   # imprime cada aba na largura de uma página
    w.sheet_properties.pageSetUpPr.fitToPage = True; w.page_setup.fitToWidth = 1; w.page_setup.fitToHeight = 0
wb.save(f'{OUT}/Conta-Doce-precificacao-confeitaria.xlsx')

# Order bump: Agenda de Encomendas
ag = Workbook(); a = ag.active; a.title = 'Encomendas'
larguras(a, {'A': 13, 'B': 18, 'C': 26, 'D': 8, 'E': 14, 'F': 14, 'G': 14, 'H': 14, 'I': 14})
cell(a, 'A1', 'Agenda de Encomendas', H1)
cell(a, 'A2', 'Guarde só o necessário: primeiro nome e o pedido. Não anote dados pessoais que você não usa.', S)
for col, t in zip('ABCDEFGHI', ['Entrega', 'Cliente', 'Pedido', 'Qtd', 'Valor total', 'Sinal pago', 'Falta receber', 'Status', 'Mês']):
    cell(a, f'{col}4', t, B)
st = DataValidation(type='list', formula1='"Orçado,Confirmado,Em produção,Entregue,Cancelado"', allow_blank=True); a.add_data_validation(st)
for r in range(5, 205):
    for col, fmt in zip('ABCDEF', ['dd/mm/yyyy', None, None, '0', R, R]): cell(a, f'{col}{r}', None, fmt=fmt, f=INP)
    cell(a, f'G{r}', f'=IF(E{r}="","",E{r}-N(F{r}))', fmt=R, f=RES)
    cell(a, f'H{r}', None, f=INP); st.add(f'H{r}')
    cell(a, f'I{r}', f'=IF(A{r}="","",TEXT(A{r},"mm/yyyy"))', f=RES)
a.conditional_formatting.add('A5:I204', FormulaRule(formula=['AND($A5<>"",$A5<=TODAY()+2,$H5<>"Entregue",$H5<>"Cancelado")'], fill=fill('FCE8B2')))
rs = ag.create_sheet('Resumo do mês'); larguras(rs, {'A': 36, 'B': 18})
cell(rs, 'A1', 'Resumo do mês', H1)
cell(rs, 'A3', 'Mês (mm/aaaa)'); cell(rs, 'B3', '=TEXT(TODAY(),"mm/yyyy")', f=INP)
for row, t, fml in [(5, 'Pedidos confirmados ou entregues', '=COUNTIFS(Encomendas!I5:I204,B3,Encomendas!H5:H204,"Confirmado")+COUNTIFS(Encomendas!I5:I204,B3,Encomendas!H5:H204,"Entregue")+COUNTIFS(Encomendas!I5:I204,B3,Encomendas!H5:H204,"Em produção")'),
                    (6, 'Faturamento do mês', '=SUMIFS(Encomendas!E5:E204,Encomendas!I5:I204,B3)-SUMIFS(Encomendas!E5:E204,Encomendas!I5:I204,B3,Encomendas!H5:H204,"Cancelado")-SUMIFS(Encomendas!E5:E204,Encomendas!I5:I204,B3,Encomendas!H5:H204,"Orçado")'),
                    (7, 'Ainda a receber', '=SUMIFS(Encomendas!G5:G204,Encomendas!I5:I204,B3,Encomendas!H5:H204,"<>Cancelado",Encomendas!H5:H204,"<>Orçado")')]:
    cell(rs, f'A{row}', t, B); cell(rs, f'B{row}', fml, B, R if row > 5 else '0', RES)
ag.save(f'{OUT}/Agenda-de-Encomendas.xlsx')
print('ok')
